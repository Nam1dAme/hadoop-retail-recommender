"""T-02/T-03: frozen field policy, byte-level TSV and the hand-calculated fixture."""
from datetime import datetime
from decimal import Decimal
from pathlib import Path
import json
import pytest
from openpyxl import Workbook
from etl.normalize import normalize,HEADERS,MAX_I64
from etl.tsv import encode,decode,line,parse_line
from etl.audit import convert

SHA='a'*64
BASE=['100001','10001','description',1,datetime(2011,8,1,12),Decimal('1.00'),1,'UK']


def workbook(path,data,extra_blank=False):
    book=Workbook();sheet=book.active;sheet.title='Online Retail'
    sheet.append(HEADERS)
    for values in data:sheet.append(values)
    if extra_blank:sheet.cell(len(data)+20,1).number_format='@'
    book.save(path);book.close()
    return path


@pytest.mark.parametrize('value',[None,'',r'\N',r'\\N',r'a\b','中文','123'])
def test_tsv_lossless_null_and_backslashes(value):
    assert decode(encode(value))==value
    assert parse_line(line([value,'a']),2)==[value,'a']


def test_tsv_rejects_controls_invalid_escapes_and_schema():
    for value in ('a\tb','a\nb','a\rb','a\x00b'):
        with pytest.raises(ValueError):encode(value)
    with pytest.raises(ValueError):decode(r'\x')
    with pytest.raises(ValueError):parse_line('a\tb\r\n',2)
    with pytest.raises(ValueError):parse_line('a\n',2)


@pytest.mark.parametrize('index,value,reason',[
 (0,None,'MISSING_INVOICE'),(0,'a'*33,'INVALID_INVOICE_LENGTH'),
 (1,None,'MISSING_STOCK_CODE'),(1,'x'*33,'INVALID_STOCK_LENGTH'),
 (3,'1.5','INVALID_QUANTITY'),(3,2**63,'INVALID_QUANTITY'),(3,0,'ZERO_QUANTITY'),
 (4,'01/02/2011','INVALID_DATETIME'),(5,'NaN','INVALID_PRICE'),(5,0,'NONPOSITIVE_PRICE'),
 (7,'x'*65,'INVALID_COUNTRY_LENGTH'),(3,MAX_I64,'AMOUNT_OVERFLOW')])
def test_rejection_priority_and_row_trace(index,value,reason):
    row=BASE.copy();row[index]=value
    if reason=='AMOUNT_OVERFLOW':row[5]=2
    result=normalize(row,SHA,2)
    assert result.record_class=='REJECTED' and result.primary_reason==reason
    assert result.fields[0]==__import__('hashlib').sha256((SHA+':2').encode()).hexdigest()


def test_rounding_customers_returns_description_and_source_time():
    values=BASE.copy();values[0]=' c123 ';values[1]='00123ab';values[2]='x\t\r\n\x00'+'y'*513
    values[5]=Decimal('1.005');values[6]='1.5'
    result=normalize(values,SHA,2)
    assert result.fields[2:4]==['C123','00123AB']
    assert len(result.fields[4])==512 and not any(c in result.fields[4] for c in '\t\r\n\x00')
    assert result.fields[7]==101 and result.fields[8] is None
    assert result.amount==-101 and result.record_class=='RETURN'
    assert set(result.warnings)=={'DESCRIPTION_TRUNCATED','PRICE_ROUNDED','INVALID_CUSTOMER_ID','CANCEL_POSITIVE_QUANTITY'}
    values=BASE.copy();values[3]=-2;values[6]=None
    result=normalize(values,SHA,3)
    assert result.record_class=='RETURN' and result.amount==-200
    assert set(result.warnings)=={'NEGATIVE_WITHOUT_CANCEL','MISSING_CUSTOMER_ID'}
    values=BASE.copy();values[4]='31/08/2011 12:34'
    assert normalize(values,SHA,4,'%d/%m/%Y %H:%M').fields[6]=='2011-08-31T12:34:00'
    values[0]=None;values[3]=0
    assert normalize(values,SHA,4).primary_reason=='MISSING_INVOICE'
    assert normalize(BASE,SHA,4,cell_error=True).primary_reason=='PARSE_ERROR'


def test_hand_calculated_fixture_and_quality_gate(tmp_path):
    specs=[(1,'100001',1,'10001',2,1),(1,'100001',1,'10002',1,3),
       (2,'100002',1,'10001',1,1),(2,'100002',1,'10003',2,2),
       (3,'100003',2,'10001',1,1),(3,'100003',2,'10002',1,3),
       (4,'C100004',1,'10002',-1,3),(5,'100005',None,'10004',1,4),
       (5,'100006',3,'10001',0,1),(6,'100007',2,'10003',1,2),(6,'100008',3,'10001',1,1)]
    data=[[invoice,stock,'item',quantity,datetime(2011,8,day,12),price,customer,'UK'] for day,invoice,customer,stock,quantity,price in specs]
    source=workbook(tmp_path/'source.xlsx',data,extra_blank=True)
    report=convert(source,tmp_path/'run',SHA)
    assert report['classification_counts']=={'SALE':9,'RETURN':1,'REJECTED':1}
    assert report['raw_rows']==11
    assert report['preflight_totals']['gross_pence']==2100 and report['preflight_totals']['return_pence']==300
    assert report['quality_gate_failures']==['REJECTION_RATE_EXCEEDED']
    assert len((tmp_path/'run/normalized.tsv').read_text().splitlines())==11
    assert json.loads((tmp_path/'run/rejected-audit.jsonl').read_text())['source_row_number']==10


def test_duplicate_retention_internal_blank_and_order_identity(tmp_path):
    data=[BASE.copy(),BASE.copy(),[None]*8,BASE.copy()]
    data[-1][4]=datetime(2011,8,2,12);data[-1][6]=2;data[-1][7]='France'
    source=workbook(tmp_path/'source.xlsx',data,extra_blank=True)
    report=convert(source,tmp_path/'run',SHA)
    assert report['raw_rows']==4 and report['metrics']['DUPLICATE_EXCESS_ROWS']==1
    assert report['conflicting_orders']==1 and 'ORDER_IDENTITY_CONFLICT' in report['quality_gate_failures']
    assert len((tmp_path/'run/normalized.tsv').read_text().splitlines())==4
    assert len((tmp_path/'run/order-conflicts.tsv').read_text().splitlines())==3


def test_null_customer_can_coexist_with_one_valid_customer(tmp_path):
    data=[BASE.copy(),BASE.copy()];data[0][6]=None
    report=convert(workbook(tmp_path/'source.xlsx',data),tmp_path/'run',SHA)
    assert report['quality_gate_passed'] and report['conflicting_orders']==0


def test_bad_header_formulas_and_aggregate_overflow(tmp_path):
    data=[BASE.copy(),BASE.copy()];data[0][0]='=SUM(1,2)'
    report=convert(workbook(tmp_path/'formula.xlsx',data),tmp_path/'formula',SHA)
    assert report['metrics']['PARSE_ERROR_ROWS']==1
    data=[BASE.copy(),BASE.copy()]
    for row in data:row[3]=10**12;row[5]=Decimal('50000.00')
    report=convert(workbook(tmp_path/'overflow.xlsx',data),tmp_path/'overflow',SHA)
    assert report['metrics']['SALE_ROWS']==2 and 'AGGREGATE_OVERFLOW' in report['quality_gate_failures']


def test_headers_multiple_sheets_and_preserving_explicit_null_strings(tmp_path):
    from etl.workbook import rows
    source=workbook(tmp_path/'source.xlsx',[BASE.copy()])
    from openpyxl import load_workbook
    book=load_workbook(source);book.create_sheet('extra');book.save(source);book.close()
    with pytest.raises(ValueError):
        with rows(source):pass
    with rows(source,'Online Retail') as (_,records):assert len(list(records))==1
    book=load_workbook(source);book['Online Retail']['A1']='WrongInvoice';book.save(source);book.close()
    with pytest.raises(ValueError):
        with rows(source,'Online Retail'):pass
    values=BASE.copy();values[2]=r'\N'
    result=normalize(values,SHA,2)
    decoded=parse_line(line(result.fields),11)
    assert decoded[4]==r'\N' and decoded[10] is None


def test_duplicate_comparison_uses_normalized_null_values(tmp_path):
    data=[BASE.copy(),BASE.copy(),BASE.copy()]
    data[0][6]=None;data[1][6]='invalid';data[2][6]='also_invalid'
    report=convert(workbook(tmp_path/'source.xlsx',data),tmp_path/'run',SHA)
    assert report['metrics']['DUPLICATE_EXCESS_ROWS']==2
    assert report['metrics']['INVALID_CUSTOMER_ID_ROWS']==2 and report['metrics']['MISSING_CUSTOMER_ID_ROWS']==1


def test_independent_scan_detects_bad_schema_and_wrong_source_row_ids(tmp_path):
    import runpy
    from orchestration.manifests import LocalStore,ManifestError
    scan=runpy.run_path(str(Path(__file__).resolve().parents[1]/'scripts/verify-w03.py'))['scan_normalized']
    values=normalize(BASE,SHA,2).fields
    source=tmp_path/'normalized.tsv';source.write_text(line(values))
    assert scan(LocalStore(tmp_path),str(source),SHA,1,tmp_path/'ok.sqlite')['source_row_ids_verified']==1
    values[0]='0'*64;source.write_text(line(values))
    with pytest.raises(ManifestError):scan(LocalStore(tmp_path),str(source),SHA,1,tmp_path/'bad-id.sqlite')
    source.write_text('only\tone\n')
    with pytest.raises(ValueError):scan(LocalStore(tmp_path),str(source),SHA,1,tmp_path/'bad-schema.sqlite')


def test_rejected_audit_keeps_all_field_errors_in_priority_order(tmp_path):
    values=BASE.copy();values[0]=None;values[3]=0;values[5]=0
    convert(workbook(tmp_path/'source.xlsx',[values]),tmp_path/'run',SHA)
    audit=json.loads((tmp_path/'run/rejected-audit.jsonl').read_text())
    assert audit['primary_reason']=='MISSING_INVOICE'
    assert audit['field_errors']==['MISSING_INVOICE','ZERO_QUANTITY','NONPOSITIVE_PRICE']
    assert audit['raw']['InvoiceNo'] is None and audit['raw']['Quantity']==0
