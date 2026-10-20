"""Streaming workbook access. Never execute formulas or silently discard internal blank rows."""
from contextlib import contextmanager
from openpyxl import load_workbook
from etl.normalize import HEADERS


@contextmanager
def rows(path, sheet=None):
    book = load_workbook(path,read_only=True,data_only=False,keep_links=False)
    try:
        if sheet is None:
            if len(book.sheetnames)!=1:
                raise ValueError('Specify a worksheet when the source has multiple sheets')
            sheet = book.sheetnames[0]
        if sheet not in book.sheetnames:
            raise ValueError('Source worksheet does not exist')
        worksheet = book[sheet]
        worksheet.reset_dimensions()  # Do not trust a stale XLSX dimension to truncate input.
        iterator = iter(worksheet.iter_rows())
        header = next(iterator,None)
        if header is None or tuple(cell.value for cell in header)!=HEADERS:
            raise ValueError('Source header must match the eight frozen UCI columns')
        def source_rows():
            blanks = 0
            for number,cells in enumerate(iterator,start=2):
                if len(cells)>8 and any(cell.value is not None for cell in cells[8:]):
                    raise ValueError('Unexpected source columns')
                values = [cell.value for cell in cells[:8]]
                values += [None]*(8-len(values))
                empty = all(value is None for value in values)
                if empty:
                    blanks += 1
                    continue
                for blank in range(number-blanks,number):
                    yield blank,[None]*8,False
                blanks = 0
                yield number,values,any(cell.data_type in ('f','e') for cell in cells[:8])
            # Trailing empty Excel rows do not count as source records.
        yield sheet,source_rows()
    finally:
        book.close()
